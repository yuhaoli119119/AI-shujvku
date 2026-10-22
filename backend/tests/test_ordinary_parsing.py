"""Isolated real PostgreSQL/API plus transport fault injection; never production."""
import asyncio
import copy
import hashlib
import importlib.util
from pathlib import Path
from uuid import uuid4

import pytest
from sqlalchemy import func, select

from app.db.models import RebuildPaperFile, WorkflowJob
from app.services.codex_web_dispatch import ThreadStatus
from app.services.ordinary_parsing import file_plan, progress_summary, reading_text
from test_rebuild_mutation_audit import audit_env, row_body


@pytest.fixture
def env(audit_env, monkeypatch):
    async def no_dispatch(**kw):
        pass
    monkeypatch.setattr('app.services.ai_extract_service._background_dispatch', no_dispatch)
    return audit_env


def start(env, **overrides):
    body = {'model': 'cn:deepseek-v4.1-flash', 'request_id': 'start-1', **overrides}
    return env['client'].post(f'/api/rebuild/papers/{env["paper"]}/ai-extract/jobs', json=body)


def cp(env, job, **overrides):
    body = {'request_id': 'cp-1', 'file_id': str(env['file']), 'stage': 'read', 'status': 'completed', 'expected_revision': 0, **overrides}
    return env['client'].put(f'/api/rebuild/papers/{env["paper"]}/ai-extract/jobs/{job}/checkpoints', json=body)


def test_one_start_prompt_bound_and_request_replay_survives_restart(env):
    response = start(env)
    assert response.status_code == 200, response.text
    job = response.json()
    assert job['file_progress']['total_files'] == 1
    assert start(env).json()['job_id'] == job['job_id']
    assert start(env, retry=True).status_code == 409
    with env['factory']() as session:
        saved = session.get(WorkflowJob, job['job_id'])
        assert len(saved.payload['dispatch_prompt']) <= 4000
        assert session.scalar(select(func.count()).select_from(WorkflowJob).where(WorkflowJob.type == 'ai_extract_rebuild')) == 1
        assert saved.progress['file_plan'][str(env['file'])]['source_sha256']


def test_checkpoint_idempotence_cas_scope_and_terminal_preserved(env):
    jid = start(env).json()['job_id']
    assert cp(env, jid).status_code == 200
    assert cp(env, jid).json()['replayed']
    assert cp(env, jid, status='failed', reason='bad').status_code == 409
    assert cp(env, jid, request_id='cp-2', stage='assets').status_code == 409
    assert cp(env, jid, request_id='cp-2', file_id=str(uuid4()), expected_revision=1).status_code == 409
    assert cp(env, jid, request_id='cp-2', stage='assets', status='partial', reason='figure unreadable', expected_revision=1).status_code == 200
    assert cp(env, jid, request_id='cp-3', status='failed', reason='bad', expected_revision=2).status_code == 409
    with env['factory']() as session:
        saved = session.get(WorkflowJob, jid)
        assert saved.progress['checkpoint_revision'] == 2
        assert saved.progress['file_plan'][str(env['file'])]['stages']['read']['status'] == 'completed'


def test_local_bad_file_keeps_good_file_and_retry_copies_completed_stage(env):
    bad = env['settings'].storage_root / 'bad.pdf'; bad.write_bytes(b'broken')
    with env['factory']() as session:
        record = RebuildPaperFile(paper_id=env['paper'], role='main', storage_path='bad.pdf', original_filename='bad.pdf', sha256=hashlib.sha256(bad.read_bytes()).hexdigest(), file_size=bad.stat().st_size)
        session.add(record); session.commit()
    first = start(env).json()
    assert first['file_progress']['total_files'] == 2
    assert any(f.get('error') for f in first['file_progress']['files'])
    assert cp(env, first['job_id']).status_code == 200
    with env['factory']() as session:
        job = session.get(WorkflowJob, first['job_id']); job.status = 'failed'
        job.progress = {**job.progress, 'native_terminal_verified': True}; session.commit()
    retry = start(env, request_id='retry-1', retry=True).json()
    assert retry['job_id'] != first['job_id']
    good = next(f for f in retry['file_progress']['files'] if f['file_id'] == str(env['file']))
    assert good['stages']['read']['status'] == 'completed'
    assert start(env, request_id='retry-1', retry=True).json()['job_id'] == retry['job_id']


def test_unknown_writer_and_failed_without_terminal_proof_never_recreated(env):
    jid = start(env).json()['job_id']
    with env['factory']() as session:
        job = session.get(WorkflowJob, jid); job.status = 'failed'; session.commit()
    assert start(env, request_id='retry', retry=True).json()['job_id'] == jid
    with env['factory']() as session:
        session.add(WorkflowJob(job_id='unknown', type='rebuild_external_extract', status='running', payload={'paper_id':str(env['paper'])}));session.commit()
    assert start(env, request_id='other').status_code == 409


def test_native_terminal_reconciles_even_no_new_values_and_keeps_partial(env, monkeypatch):
    from app.services.ai_extract_service import sync_job_status
    jid = start(env).json()['job_id']
    async def status(*a, **kw):
        return ThreadStatus('completed', 'complete', None, 'partial ordinary results', thread_id='owned', turn_id='turn', item_count=1)
    monkeypatch.setattr('app.services.codex_web_dispatch.read_thread_status', status)
    with env['factory']() as session:
        job=session.get(WorkflowJob,jid);job.status='running';job.progress={**job.progress,'thread_id':'owned','turn_id':'turn'};session.commit()
    with env['factory']() as session:
        job=asyncio.run(sync_job_status(session, session.get(WorkflowJob,jid), env['settings']));session.commit()
        assert job.status=='completed' and job.progress['native_terminal_verified']
        assert job.progress['content_status']=='partial'
        assert job.result['file_ids']==[]
        assert job.result['scientific_acceptance']=='pending'


def test_active_goal_failed_turn_and_foreign_turn_do_not_release_writer(env, monkeypatch):
    from app.services.ai_extract_service import sync_job_status
    jid=start(env).json()['job_id']
    for observed in [ThreadStatus('failed','active',None,None,thread_id='owned',turn_id='turn'), ThreadStatus('completed','complete',None,None,thread_id='foreign',turn_id='turn')]:
        async def status(*a, **kw): return observed
        monkeypatch.setattr('app.services.codex_web_dispatch.read_thread_status',status)
        with env['factory']() as session:
            job=session.get(WorkflowJob,jid);job.status='running';job.progress={**job.progress,'thread_id':'owned','turn_id':'turn'};session.commit()
        with env['factory']() as session:
            job=asyncio.run(sync_job_status(session,session.get(WorkflowJob,jid),env['settings']));session.commit()
            assert job.status=='running' and not job.progress.get('native_terminal_verified')


def test_docx_reading_locators_match_provenance_validation(env):
    from docx import Document
    from app.services.rebuild_file_preview import validate_ooxml_identity
    path=env['settings'].storage_root/'source.docx'
    document=Document();document.add_paragraph('Original paragraph');table=document.add_table(rows=2,cols=2);table.cell(0,0).text='energy';table.cell(1,1).text='0.12';document.save(path)
    with env['factory']() as session:
        file=RebuildPaperFile(paper_id=env['paper'],role='si',storage_path='source.docx',original_filename='source.docx',sha256=hashlib.sha256(path.read_bytes()).hexdigest(),file_size=path.stat().st_size)
        session.add(file);session.flush()
        result=reading_text(session,env['paper'],file.id,env['settings'],0,1)
        assert result['page_number'] is None and result['has_more']
        data={'page_basis':'ooxml','source_sha256':file.sha256,'page_number':None,'ooxml_locator':result['items'][0]['ooxml_locator']}
        assert validate_ooxml_identity(file,env['settings'],data)['page_basis']=='ooxml'
        result=reading_text(session,env['paper'],file.id,env['settings'],1,1)
        data['ooxml_locator']=result['items'][0]['cell_locators'][1][1]
        assert validate_ooxml_identity(file,env['settings'],data)['ooxml_locator']==data['ooxml_locator']


def test_recovery_replays_only_never_sent_intent(env, monkeypatch):
    from app.services.ordinary_parsing import reconcile_once
    jid=start(env).json()['job_id'];calls=[]
    async def dispatch(**kw): calls.append(kw['job_id'])
    monkeypatch.setattr('app.services.ai_extract_service._background_dispatch',dispatch)
    asyncio.run(reconcile_once(env['settings']))
    assert calls==[jid]
    with env['factory']() as session:
        job=session.get(WorkflowJob,jid);job.progress={**job.progress,'dispatch_started':True};session.commit()
    calls.clear();asyncio.run(reconcile_once(env['settings']))
    assert calls==[]


@pytest.mark.no_test_database
def test_client_429_backoff_and_lost_ack_reuses_exact_bytes(tmp_path):
    import io
    import urllib.error
    import importlib.util
    spec=importlib.util.spec_from_file_location('ordinary_client',Path(__file__).resolve().parents[2]/'scripts/ordinary_parse_client.py')
    module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
    class Response(io.BytesIO):
        def __enter__(self):return self
        def __exit__(self,*a):self.close()
    class Opener:
        def __init__(self):self.calls=[]
        def open(self,req,timeout):
            self.calls.append(req.data)
            if len(self.calls)<4:raise urllib.error.HTTPError(req.full_url,429,'rate limited',{'Retry-After':'1'},io.BytesIO(b'limited'))
            return Response(b'{"complete":true}')
    opener=Opener();delays=[];client=module.WorkbenchClient('http://fixture',opener=opener,sleep=delays.append)
    result=module.durable_write(client,tmp_path,'/api/rebuild/papers/fixture/assets/import',{'request_id':'same','assets':[]})
    assert result['complete'] and delays==[1,1,1] and len(set(opener.calls))==1
    assert module.durable_write(client,tmp_path,'/api/rebuild/papers/fixture/assets/import',{'request_id':'same','assets':[]})==result
    assert len(opener.calls)==4
    class LostAck:
        def __init__(self):self.calls=[]
        def open(self,req,timeout):
            self.calls.append(req.data)
            if len(self.calls)<3:raise urllib.error.URLError('connection lost after commit')
            return Response(b'{"replayed":true}')
    opener=LostAck();client=module.WorkbenchClient('http://fixture',opener=opener,sleep=lambda _:None)
    assert client.request('/api/rebuild/rows/import',b'exact',idempotent=True)['replayed']
    assert opener.calls==[b'exact']*3


def test_bad_source_can_record_failure_without_losing_siblings(env):
    jid = start(env).json()['job_id']
    with env['factory']() as session:
        record = session.get(RebuildPaperFile, env['file'])
        source = env['settings'].storage_root / record.storage_path
        source.write_bytes(b'synthetic corruption')
    assert cp(env,jid,status='completed').status_code==409
    assert cp(env,jid,status='failed',reason='source checksum changed').status_code==200


def test_subset_retry_preserves_earlier_completed_files(env):
    from app.services.ordinary_parsing import saved_file_plan
    jid=start(env).json()['job_id']
    assert cp(env,jid).status_code==200
    with env['factory']() as session:
        old=session.get(WorkflowJob,jid);old.status='completed'
        session.add(WorkflowJob(job_id='aiext-subset',type='ai_extract_rebuild',status='completed',payload={'paper_id':str(env['paper']),'parse_contract':'ordinary_parse_v1'},progress={'file_plan':{}}));session.commit()
    with env['factory']() as session:
        assert saved_file_plan(session,env['paper'])[str(env['file'])]['stages']['read']['status']=='completed'


@pytest.mark.no_test_database
def test_journal_publish_interruption_leaves_no_partial_official_json(tmp_path, monkeypatch):
    spec=importlib.util.spec_from_file_location('ordinary_client_atomic',Path(__file__).resolve().parents[2]/'scripts/ordinary_parse_client.py')
    module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
    fsync=module.os.fsync
    def interrupted(fd): raise OSError('synthetic process interruption before publish')
    monkeypatch.setattr(module.os,'fsync',interrupted)
    target=tmp_path/'official.json'
    with pytest.raises(OSError):module.publish_json(target,{'request':'exact'})
    assert not target.exists()
    monkeypatch.setattr(module.os,'fsync',fsync)
    assert module.publish_json(target,{'request':'exact'})
    assert not module.publish_json(target,{'request':'exact'})


@pytest.mark.no_test_database
def test_client_rate_limit_budget_exhausts_after_three_retries():
    import io, urllib.error
    spec=importlib.util.spec_from_file_location('ordinary_client_budget',Path(__file__).resolve().parents[2]/'scripts/ordinary_parse_client.py')
    module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
    calls=[];delays=[]
    class Opener:
        def open(self,req,timeout):
            calls.append(req.data)
            raise urllib.error.HTTPError(req.full_url,429,'rate limited',{},io.BytesIO(b'limited'))
    client=module.WorkbenchClient('http://fixture',opener=Opener(),sleep=delays.append)
    with pytest.raises(RuntimeError,match='HTTP 429'):
        client.request('/api/rebuild/rows/import',b'body',idempotent=True)
    assert calls==[b'body']*4 and delays==[5,10,20]
