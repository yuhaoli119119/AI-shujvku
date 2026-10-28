"""Prepare U137, U141 and U142 source-bound row payloads; no I/O writes."""
from __future__ import annotations
import csv
import io
import math
import re
from decimal import Decimal, InvalidOperation
from zipfile import ZipFile
from pathlib import Path
import docx
import openpyxl


def _cell(index, label, value, file_record, sheet, row, *, unit=None, basis=None, original=None):
    locator = {'sha256': file_record['sha256'], 'sheet': sheet, 'row': row, 'column': label}
    if original is not None:
        locator['original_value'] = original
    source = {'source_kind': 'dataset', 'dataset_file_id': file_record['id'], 'dataset_locator': locator}
    if basis:
        source['estimate_basis'] = basis
    raw = str(value)
    numeric = None
    if re.fullmatch(r'[+\-−]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[eE][+\-]?[0-9]{1,6})?', raw):
        try:
            exact = Decimal(raw.replace('−', '-')); candidate = float(exact)
            if math.isfinite(candidate) and not (candidate == 0 and exact != 0): numeric = candidate
        except (ValueError, InvalidOperation, OverflowError):
            pass
    return {'field_name': 'dataset_col_%03d' % index, 'raw_value': raw, 'numeric_value': numeric, 'unit': unit,
            'value_type': 'derived' if basis else 'explicit', 'sources': [source]}


def _row(file_sha, member, position, material, values, kind, notes):
    identity = f'dataset:{file_sha}:{member}:{position}'
    return {'row_key': identity, 'reaction': 'SRR', 'material': material, 'active_site_type': 'other',
            'data_type': 'dft', 'condition': {'identity_context': identity, 'dataset_kind': kind},
            'notes': notes, 'values': values}


def build(root: Path, files: dict[str, dict]):
    """files maps registered names to {id, sha256}, after complete hash verification."""
    out = {'U137': [], 'U141': [], 'U142': []}
    zip_name = 'u137-github-v1.1.0.zip'
    with ZipFile(root / 'U137' / zip_name) as archive:
        for filename in ('data_2A.csv', 'data_3A.csv'):
            matches = [n for n in archive.namelist() if n.endswith('/data/' + filename)]
            if len(matches) != 1: raise ValueError('ambiguous archive member')
            member = matches[0]
            with archive.open(member) as binary:
                reader = csv.DictReader(io.TextIOWrapper(binary, encoding='utf-8-sig', newline=''))
                for position, data in enumerate(reader, 2):
                    if not data.get('formula'): raise ValueError('missing U137 formula')
                    values = [_cell(i, name, data[name], files[zip_name], member, position,
                                    unit='eV' if name == 'target' else None)
                              for i, name in enumerate(reader.fieldnames, 1) if data[name] not in ('', None)]
                    out['U137'].append(_row(files[zip_name]['sha256'], member, position, data['formula'], values,
                                            'dft_adsorption_descriptors', 'target 的 eV 单位依据论文；原 CSV 本身未标单位。'))
    workbook_name = 'u141-ml4lis-data.xlsx'
    csv_name = 'u141_ML4LiS_descriptor_dataset_95rows.csv'
    workbook = openpyxl.load_workbook(root / 'U141' / workbook_name, read_only=True, data_only=True)
    sheet = workbook['Sheet1']; excel = list(sheet.values); headers = excel[0]
    with (root / 'U141' / csv_name).open(encoding='utf-8-sig', newline='') as inp:
        for data in csv.DictReader(inp):
            source_row = int(data['source_row']); actual = excel[source_row]
            if any(float(data[str(name).strip()]) != float(value) for name, value in zip(headers, actual)):
                raise ValueError(f'U141 source mismatch at row {source_row}')
            values = [_cell(i, name, value, files[workbook_name], 'Sheet1', source_row + 1,
                            unit='eV' if name == 'G' else None)
                      for i, (name, value) in enumerate(zip(headers, actual), 1) if value is not None]
            values.append(_cell(10, 'adsorbate', data['adsorbate'], files[csv_name], 'CSV', source_row + 1,
                                basis='整理后的 CSV 指定吸附物分类；原工作簿未提供该字段。'))
            out['U141'].append(_row(files[workbook_name]['sha256'], 'Sheet1', source_row + 1,
                                    '催化剂身份未报告', values, 'descriptor_dataset',
                                    '描述符数据集，部分可复现；缺催化剂身份，脚本期望13特征，公开表仅8个描述符和G；许可证待确认。'))
    workbook.close()
    docx_name = 'u142-si.docx'; xlsx_name = 'u142_SI_Tables_S1-S7.xlsx'
    document = docx.Document(root / 'U142' / docx_name)
    captions = {i: next((p.text.strip() for p in document.paragraphs if p.text.strip().startswith(f'Table S{i}.')), None) for i in range(1, 8)}
    if any(not value for value in captions.values()): raise ValueError('U142 caption missing')
    workbook = openpyxl.load_workbook(root / 'U142' / xlsx_name, read_only=True, data_only=True)
    worksheets = [s for s in workbook if s.title != 'Metadata']
    if len(document.tables) != 7 or len(worksheets) != 7: raise ValueError('U142 table count mismatch')
    for index, (table, ws) in enumerate(zip(document.tables, worksheets), 1):
        rows = list(ws.values); headers = rows[0]
        if len(table.rows) != len(rows): raise ValueError('U142 row mismatch')
        for position, (docrow, extracted) in enumerate(zip(list(table.rows)[1:], rows[1:]), 2):
            originals = [cell.text.strip() for cell in docrow.cells]
            for old, new in zip(originals, extracted):
                if old.replace('−','-') != str(new if new is not None else '').strip().replace('−','-'):
                    raise ValueError(f'U142 raw value mismatch S{index} row {position}')
            values = []
            for col, value in enumerate(extracted, 1):
                if value is None or str(value) == '': continue
                header = str(headers[col-1]); original_col = col; basis = None
                if index == 2 and col > len(originals):
                    original_col = {8:3, 9:3, 10:4, 11:4, 12:5, 13:5}[col]
                    basis = '由原表斜线分隔的配对数值按顺序拆列；原单元格值保存在来源中。'
                    parts = originals[original_col-1].replace('−', '-').split('/')
                    part_index = (col - 8) % 2
                    if len(parts) != 2 or float(parts[part_index].strip()) != float(value):
                        raise ValueError(f'U142 split mismatch S{index} row {position} column {col}')
                original_header = table.rows[0].cells[original_col-1].text.strip()
                label = f'column_{original_col}' + (f' ({original_header})' if original_header else '')
                values.append(_cell(col, label, value, files[docx_name], f'Table S{index}', position,
                                    unit='eV' if header.endswith('_eV') else None, basis=basis,
                                    original=originals[original_col-1]))
            out['U142'].append(_row(files[docx_name]['sha256'], f'Table S{index}', position,
                                    originals[0] or '材料身份未报告', values, 'si_table',
                                    '来自原 DOCX 表；不能可靠识别的公式表头使用中性列名。 ' + captions[index]))
    workbook.close()
    if {k:len(v) for k,v in out.items()} != {'U137':51,'U141':95,'U142':113}:
        raise ValueError('unexpected row count')
    return out
